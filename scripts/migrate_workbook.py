"""Read-only, one-time migration. No worksheet is edited or recalculated.

Usage: python scripts/migrate_workbook.py source.xlsx output.json
Requires openpyxl. Source values + formulas are retained in the JSON backup.
"""
import datetime as dt
import json
import math
from pathlib import Path
import re
import sys
from collections import Counter, defaultdict

SHEETS = ['food', 'recipes', 'me', 'Daily Tracker', 'past_log']
MACROS = ['carbs', 'protein', 'fat']

def valid(values):
    return all(isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) and v >= 0 for v in values)

def macro(values):
    return dict(zip(MACROS, values))

def migrate(source):
    import openpyxl
    formulas = openpyxl.load_workbook(source, data_only=False)
    values = openpyxl.load_workbook(source, data_only=True)
    now = dt.datetime.now(dt.timezone.utc).isoformat()
    state = dict(version=1, foods=[], entries=[], presets=[], days={}, weights={}, profile={}, imports=[], notes=[], raw={})
    notes = state['notes']
    for name in SHEETS:
        sheet = formulas[name]
        state['raw'][name] = {c.coordinate: {'formulaOrValue': c.value, 'cachedValue': values[name][c.coordinate].value} for row in sheet for c in row if c.value is not None}
    for row in values['food'].iter_rows(min_row=2):
        r = row[0].row
        name, unit, c, p, f = [cell.value for cell in row[:5]]
        if not str(name or '').strip():
            continue
        if not unit or not valid([c, p, f]):
            notes.append(f'food row {r}: incomplete nutrition or unit. Original preserved; not imported as a selectable food.')
            continue
        unit = re.sub(r'^1\s*', '', str(unit).strip()) if str(unit).strip() != '100g' else '100g'
        state['foods'].append(dict(id=f'food-{r}', name=str(name).strip(), unit=unit, macros=macro([c,p,f]), type='food', source=f'food row {r}', createdAt=now, updatedAt=now))
        if unit == '100g' and sum([c,p,f]) > 100:
            notes.append(f'food row {r} ({name}): macros exceed 100g per 100g; imported exactly as supplied, please review.')
    duplicates = [name for name,count in Counter(f['name'].casefold() for f in state['foods']).items() if count > 1]
    if duplicates:
        notes.append('Duplicate food names retained as separate source rows: '+', '.join(duplicates)+'. Recipe lookup follows Excel’s first exact match.')
    lookup = {}
    for food in state['foods']:
        lookup.setdefault(food['name'].casefold(), food)
    sheet = values['recipes']
    starts = [r for r in range(2,sheet.max_row+1) if sheet.cell(r,1).value and str(sheet.cell(r,1).value).strip()]
    for ix,start in enumerate(starts):
        end=starts[ix+1] if ix+1<len(starts) else sheet.max_row+1
        name=str(sheet.cell(start,1).value).strip()
        finish=next((r for r in range(start,end) if sheet.cell(r,3).value=='actual weight'),None)
        reason=''
        if finish is None or not valid([sheet.cell(finish,4).value]) or sheet.cell(finish,4).value<=0:
            reason='missing explicit actual finished weight'
        ingredients=[]
        if not reason:
            for r in range(start,finish):
                ingredient=sheet.cell(r,2).value
                if not ingredient or not str(ingredient).strip():
                    continue
                f=lookup.get(str(ingredient).strip().casefold())
                quantity=sheet.cell(r,4).value
                if not f or not valid([quantity]) or quantity<=0:
                    reason=f'ingredient at row {r} has a missing quantity or food reference';break
                ingredients.append(dict(id=f'ingredient-{r}',foodId=f['id'],name=f['name'],unit=f['unit'],macros=f['macros'].copy(),quantity=quantity))
            weight=sheet.cell(finish,4).value
            totals={k:sum(i['macros'][k]*i['quantity']/(100 if i['unit']=='100g' else 1) for i in ingredients) for k in MACROS}
            per={k:totals[k]/weight*100 for k in MACROS}
            cached=[sheet.cell(finish,c).value for c in [5,6,7]]
            if not ingredients or not valid(cached) or any(abs(per[k]-v)>1e-7 for k,v in zip(MACROS,cached)):
                reason=reason or 'cached totals differ from ingredient calculation'
        if reason:
            notes.append(f'recipes rows {start}–{end-1} ({name}): {reason}. Preserved in original source data; not guessed. Existing food-sheet entries remain available.')
        else:
            state['foods'].append(dict(id=f'recipe-{start}',name=name,unit='100g',type='recipe',ingredients=ingredients,finishedWeight=weight,macros=per,source=f'recipes row {start}',createdAt=now,updatedAt=now))
            if finish<end-1 and any(sheet.cell(r,2).value for r in range(finish+1,end)):
                notes.append(f'recipes rows {finish+1}–{end-1}: extra rows after finished weight were preserved separately and not added to {name}.')
    me=values['me']
    for ident,name,row in [('standard','Standard',20),('high','High-carb / Cheat Day',21),('custom','Custom',20)]:
        state['presets'].append(dict(id=ident,name=name,**macro([me.cell(row,c).value for c in [15,16,17]])))
    state['profile']=dict(bmr=me['A3'].value,weight=me['B3'].value,currentFat=me['C3'].value,targetFat=me['D3'].value,days=me['G3'].value,minutes=me['B5'].value,intensity={'低':'Low','中':'Medium','高':'High'}[me['A5'].value],activity={'低':'Low','中':'Medium','高':'High'}[me['A7'].value],cardio=me['A9'].value,split=macro([me.cell(23,c).value*100 for c in [1,2,3]]))
    notes.append('Active presets use me!O20:Q21. Daily Tracker high-carb fat is 30g, whereas me uses 20g; historical comparisons retain each sheet’s targets.')
    notes.append('The undated meal entries in me are preserved in the original source data, not assigned to today.')
    for name in ['Daily Tracker','past_log']:
        sheet=values[name]
        targets={'standard':macro([sheet.cell(1,c).value for c in [3,4,5]]),'high':macro([sheet.cell(1,c).value for c in [8,9,10]])}
        for r in range(3,sheet.max_row+1):
            if str(sheet.cell(r,1).value or '').strip().casefold()=='goal change':
                for ident,cols in [('standard',[3,4,5]),('high',[8,9,10])]:
                    vals=[sheet.cell(r,c).value for c in cols]
                    if valid(vals):targets[ident]=macro(vals)
                notes.append(f'{name} row {r}: Goal Change row updates subsequent targets, not consumed food.')
                continue
            for ident,col,cols in [('standard',2,[3,4,5]),('high',7,[8,9,10])]:
                date=sheet.cell(r,col).value; vals=[sheet.cell(r,c).value for c in cols]
                if not any(v is not None for v in vals):continue
                if not isinstance(date,(dt.date,dt.datetime)) or not valid(vals):
                    notes.append(f'{name} row {r} ({ident}): undated or non-numeric totals preserved, excluded from daily sums.');continue
                day=date.strftime('%Y-%m-%d')
                goal=dict(id=f'import-{ident}',name='Standard (imported)' if ident=='standard' else 'High-carb (imported)',**targets[ident])
                state['imports'].append(dict(id=f'{name}-{r}-{ident}',date=day,macros=macro(vals),goal=goal,source=f'{name} row {r} ({ident})'))
            date=sheet.cell(r,2).value
            weight=sheet.cell(r,12).value if name=='Daily Tracker' else None
            if isinstance(date,(dt.date,dt.datetime)) and valid([weight]) and weight>0:
                state['weights'][date.strftime('%Y-%m-%d')]=weight*0.45359237
    groups=defaultdict(list)
    for record in state['imports']:groups[record['date']].append(record)
    for date,records in groups.items():
        if len(records)>1:
            for record in records:record['conflict']=True
            notes.append(f'{date}: {len(records)} conflicting daily totals retained. Select one on that date; no double-counting or automatic choice.')
        else:state['days'][date]={'goal':records[0]['goal']}
    notes.append('Daily Tracker weight is labelled lb. Imported weights are converted to kg using 0.45359237; original pounds are preserved.')
    return state

if __name__=='__main__':
    result=migrate(sys.argv[1]);out=Path(sys.argv[2]);out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str))
    print(json.dumps({'foods':sum(f['type']=='food' for f in result['foods']),'recipes':sum(f['type']=='recipe' for f in result['foods']),'historicalRecords':len(result['imports']),'resolvedDates':len(result['days']),'conflictingDates':len(set(i['date'] for i in result['imports'] if i.get('conflict'))),'notes':len(result['notes'])}))
